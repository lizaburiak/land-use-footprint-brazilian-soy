"""
Combine per-year input files into single multi-year files.
Keeps the original per-year files intact, creates combined versions alongside them.

Files combined:
  1. EXP_{YEAR}_MUN_COMEX.csv -> EXP_ALL_MUN_COMEX.csv  (already have year column CO_ANO)
  2. IMP_{YEAR}_MUN_COMEX.csv -> IMP_ALL_MUN_COMEX.csv  (already have year column CO_ANO)
  3. Production_tabela1612_IBGE_{YEAR}.csv -> Production_tabela1612_IBGE_ALL.csv
  4. Population_tabela6579_IBGE_{YEAR}.csv -> Population_tabela6579_IBGE_ALL.csv
  5. Livestock_{YEAR}_tabela3939_IBGE.csv -> Livestock_ALL_tabela3939_IBGE.csv
  6. MilkCows_{YEAR}_tabela94_IBGE.csv -> MilkCows_ALL_tabela94_IBGE.csv
  7. FAOSTAT_tradematrix_BRAsoy_{YEAR}.csv -> FAOSTAT_tradematrix_BRAsoy_ALL.csv
  8. CBS_SOY_{YEAR}_FAO.xlsx -> CBS_SOY_ALL_FAO.xlsx (multi-sheet: one sheet per year)
"""

import os
import csv
import glob
from pathlib import Path

BASE_DIR = Path(__file__).parent
INPUTS_DIR = BASE_DIR / "inputs"


def log(msg):
    print(f"[INFO] {msg}")


# ============================================================================
# 1 & 2. COMEX Trade Data (exports and imports)
# ============================================================================
# These already have CO_ANO (year) as the first column, so we just concatenate.

def combine_comex():
    for prefix, label in [("EXP", "Exports"), ("IMP", "Imports")]:
        log(f"Combining COMEX {label}...")
        files = sorted(glob.glob(str(INPUTS_DIR / "00" / f"{prefix}_*_MUN_COMEX.csv")))
        if not files:
            log(f"  No {prefix} files found, skipping")
            continue

        dest = INPUTS_DIR / "00" / f"{prefix}_ALL_MUN_COMEX.csv"
        header_written = False
        total_rows = 0

        with open(dest, "w", encoding="utf-8") as out:
            for f in files:
                with open(f, "r", encoding="utf-8") as inp:
                    header = inp.readline()
                    if not header_written:
                        out.write(header)
                        header_written = True
                    for line in inp:
                        out.write(line)
                        total_rows += 1

        log(f"  Combined {len(files)} files -> {dest.name} ({total_rows:,} rows)")


# ============================================================================
# 3. Production (SIDRA 1612)
# ============================================================================
# Original format: 2 header lines + data. Downloaded format: 3 header lines + data.
# We unify to a simple CSV: Cód.,Município,Produto,Ano,Área plantada,Área colhida,Quantidade produzida

def combine_production():
    log("Combining Production (SIDRA 1612)...")

    # Include both the original 2013 file and downloaded year files
    files_new = sorted(glob.glob(str(INPUTS_DIR / "00" / "Production_tabela1612_IBGE_*.csv")))
    file_orig = INPUTS_DIR / "00" / "Production_tabela1612_IBGE.csv"

    all_files = []
    if file_orig.exists():
        all_files.append(("orig", str(file_orig)))
    for f in files_new:
        all_files.append(("new", f))

    if not all_files:
        log("  No production files found, skipping")
        return

    dest = INPUTS_DIR / "00" / "Production_tabela1612_IBGE_ALL.csv"
    header = '"Cód.","Município","Produto das lavouras temporárias","Ano","Área plantada (Hectares)","Área colhida (Hectares)","Quantidade produzida (Toneladas)"\n'

    with open(dest, "w", encoding="utf-8") as out:
        out.write(header)
        total_rows = 0

        for ftype, fpath in all_files:
            with open(fpath, "r", encoding="utf-8-sig") as inp:
                lines = inp.readlines()
                # Skip header lines (2 for original, 3 for downloaded)
                skip = 3 if ftype == "new" else 2
                for line in lines[skip:]:
                    line = line.strip()
                    if line and not line.startswith('"Fonte') and not line.startswith('"Nota'):
                        out.write(line + "\n")
                        total_rows += 1

    log(f"  Combined {len(all_files)} files -> {dest.name} ({total_rows:,} rows)")


# ============================================================================
# 4. Population (SIDRA 6579)
# ============================================================================
# Format: 1 header line + data rows
# "Cód.","Município","Ano","Variável",""

def combine_population():
    log("Combining Population (SIDRA 6579)...")

    files_new = sorted(glob.glob(str(INPUTS_DIR / "00" / "Population_tabela6579_IBGE_*.csv")))
    file_orig = INPUTS_DIR / "00" / "Population_tabela6579_IBGE.csv"

    all_files = []
    if file_orig.exists():
        all_files.append(("orig", str(file_orig)))
    for f in files_new:
        all_files.append(("new", f))

    if not all_files:
        log("  No population files found, skipping")
        return

    dest = INPUTS_DIR / "00" / "Population_tabela6579_IBGE_ALL.csv"
    header = '"Cód.","Município","Ano","Variável",""\n'

    with open(dest, "w", encoding="utf-8") as out:
        out.write(header)
        total_rows = 0

        for ftype, fpath in all_files:
            with open(fpath, "r", encoding="utf-8-sig") as inp:
                lines = inp.readlines()
                # Skip 1 header line for original, 2 for new (title + header)
                skip = 2 if ftype == "new" else 1
                for line in lines[skip:]:
                    line = line.strip()
                    if line and not line.startswith('"Fonte') and not line.startswith('"Nota'):
                        out.write(line + "\n")
                        total_rows += 1

    log(f"  Combined {len(all_files)} files -> {dest.name} ({total_rows:,} rows)")


# ============================================================================
# 5. Livestock (SIDRA 3939)
# ============================================================================
# Format: semicolon-separated, 4-5 header lines, then data rows
# "MU";"Cód.";"Município";"Bovino";"Bubalino";...

def combine_livestock():
    log("Combining Livestock (SIDRA 3939)...")

    files = sorted(glob.glob(str(INPUTS_DIR / "00" / "Livestock_*_tabela3939_IBGE.csv")))
    if not files:
        log("  No livestock files found, skipping")
        return

    dest = INPUTS_DIR / "00" / "Livestock_ALL_tabela3939_IBGE.csv"
    col_names = ["Bovino", "Bubalino", "Equino", "Suíno - total",
                 "Suíno - matrizes de suínos", "Caprino", "Ovino",
                 "Galináceos - total", "Galináceos - galinhas", "Codornas"]
    header = '"Nível";"Cód.";"Município";"Ano";' + ";".join([f'"{c}"' for c in col_names]) + "\n"

    with open(dest, "w", encoding="utf-8") as out:
        out.write(header)
        total_rows = 0

        for fpath in files:
            # Extract year from filename
            fname = os.path.basename(fpath)
            year = fname.split("_")[1]

            with open(fpath, "r", encoding="utf-8-sig") as inp:
                lines = inp.readlines()
                # Find data lines (start with "MU")
                for line in lines:
                    line = line.strip()
                    if line.startswith('"MU"'):
                        # Insert year after municipality: "MU";"code";"name";"year";values...
                        parts = line.split(";", 3)
                        out.write(f'{parts[0]};{parts[1]};{parts[2]};"{year}";{parts[3]}\n')
                        total_rows += 1

    log(f"  Combined {len(files)} files -> {dest.name} ({total_rows:,} rows)")


# ============================================================================
# 6. Milked Cows (SIDRA 94)
# ============================================================================
# Format: semicolon-separated, 3-4 header lines, then data
# "MU";"code";"name";"value"

def combine_milkcows():
    log("Combining Milked Cows (SIDRA 94)...")

    files = sorted(glob.glob(str(INPUTS_DIR / "00" / "MilkCows_*_tabela94_IBGE.csv")))
    if not files:
        log("  No milkcows files found, skipping")
        return

    dest = INPUTS_DIR / "00" / "MilkCows_ALL_tabela94_IBGE.csv"
    header = '"Nível";"Cód.";"Município";"Ano";"Vacas ordenhadas"\n'

    with open(dest, "w", encoding="utf-8") as out:
        out.write(header)
        total_rows = 0

        for fpath in files:
            fname = os.path.basename(fpath)
            year = fname.split("_")[1]

            with open(fpath, "r", encoding="utf-8-sig") as inp:
                lines = inp.readlines()
                for line in lines:
                    line = line.strip()
                    if line.startswith('"MU"'):
                        # "MU";"code";"name";"value" -> add year
                        parts = line.split(";", 3)
                        out.write(f'{parts[0]};{parts[1]};{parts[2]};"{year}";{parts[3]}\n')
                        total_rows += 1

    log(f"  Combined {len(files)} files -> {dest.name} ({total_rows:,} rows)")


# ============================================================================
# 7. FAOSTAT Trade Matrix
# ============================================================================
# Standard CSV with Year column already present. Just concatenate.

def combine_faostat_trade():
    log("Combining FAOSTAT Trade Matrix...")

    files = sorted(glob.glob(str(INPUTS_DIR / "04" / "FAOSTAT_tradematrix_BRAsoy_*.csv")))
    # Also include original 2013 file
    file_orig = INPUTS_DIR / "04" / "FABIO" / "FAOSTAT_tradematrix_BRAsoy.csv"

    all_files = []
    if file_orig.exists():
        all_files.append(str(file_orig))
    all_files.extend(files)

    if not all_files:
        log("  No trade matrix files found, skipping")
        return

    dest = INPUTS_DIR / "04" / "FAOSTAT_tradematrix_BRAsoy_ALL.csv"
    header_written = False
    total_rows = 0

    with open(dest, "w", encoding="utf-8") as out:
        for f in all_files:
            with open(f, "r", encoding="utf-8-sig") as inp:
                header = inp.readline()
                if not header_written:
                    out.write(header)
                    header_written = True
                for line in inp:
                    out.write(line)
                    total_rows += 1

    log(f"  Combined {len(all_files)} files -> {dest.name} ({total_rows:,} rows)")


# ============================================================================
# 8. FAO CBS (Food Balance Sheets) -> multi-sheet Excel
# ============================================================================

def combine_fao_cbs():
    log("Combining FAO CBS (Food Balance Sheets)...")

    try:
        import openpyxl
    except ImportError:
        log("  openpyxl not available, skipping CBS combination")
        return

    files = sorted(glob.glob(str(INPUTS_DIR / "00" / "CBS_SOY_*_FAO.xlsx")))
    if not files:
        log("  No CBS files found, skipping")
        return

    dest = INPUTS_DIR / "00" / "CBS_SOY_ALL_FAO.xlsx"
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)  # Remove default sheet

    for fpath in files:
        fname = os.path.basename(fpath)
        year = fname.replace("CBS_SOY_", "").replace("_FAO.xlsx", "")

        wb_in = openpyxl.load_workbook(fpath)
        ws_in = wb_in.active
        ws_out = wb_out.create_sheet(title=str(year))

        for row in ws_in.iter_rows(values_only=True):
            ws_out.append(list(row))

    wb_out.save(dest)
    log(f"  Combined {len(files)} files -> {dest.name} ({len(wb_out.sheetnames)} sheets: {wb_out.sheetnames[0]}-{wb_out.sheetnames[-1]})")


# ============================================================================
# MAIN
# ============================================================================

if __name__ == "__main__":
    log("=" * 60)
    log("Combining per-year files into multi-year files")
    log("=" * 60)
    log("")

    combine_comex()
    log("")
    combine_production()
    log("")
    combine_population()
    log("")
    combine_livestock()
    log("")
    combine_milkcows()
    log("")
    combine_faostat_trade()
    log("")
    combine_fao_cbs()

    log("")
    log("=" * 60)
    log("DONE - Combined files created alongside originals")
    log("=" * 60)
    log("")
    log("Combined files:")
    log("  inputs/00/EXP_ALL_MUN_COMEX.csv")
    log("  inputs/00/IMP_ALL_MUN_COMEX.csv")
    log("  inputs/00/Production_tabela1612_IBGE_ALL.csv")
    log("  inputs/00/Population_tabela6579_IBGE_ALL.csv")
    log("  inputs/00/Livestock_ALL_tabela3939_IBGE.csv")
    log("  inputs/00/MilkCows_ALL_tabela94_IBGE.csv")
    log("  inputs/04/FAOSTAT_tradematrix_BRAsoy_ALL.csv")
    log("  inputs/00/CBS_SOY_ALL_FAO.xlsx")
